from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation


# =========================================================
# LISTES DES VALEURS POSSIBLES
# =========================================================

LIBRAIRIES = [
    "CAMP DE REVEIL ADN",
    "CAMP KHAYIL",
    "EDMONTON",
    "EJP",
    "ELGO",
    "LAVAL",
    "LONGUEUIL",
    "MONCTON",
    "MONTREAL",
    "OTTAWA",
    "QUEBEC",
    "SHAWINIGAN",
    "SHERBROOKE",
    "TIMMINS",
    "TORONTO",
    "TROIS RIVIERES",
    "VICTORIAVILLE",
]

ARTICLES = [
    "10 Questions essentielles pour trouver sa mission de vie",
    "1000 micro-églises",
    "12 clés pour prendre son envol professionnel",
    "12 fondations pour bâtir une année de succès et de foi",
    "15 défis pour un leadership performant",
    "21 Pièges à déjouer dans le. choix du conjoint pour un mariage heureux",
    "21 pensées pour sortir des prisons intérieures",
    "21 pièges à déjouer dans le choix du conjoint pour un mariage heureux",
    "3 clés maitresse pour parvenir au sommet et briller",
    "365 jours",
    "365 jours de victoire",
    "4 secrets d'un mariage réussi",
    "40 jours pour rendre votre mariage inebranlable",
    "50 Conseils pratiques pour une Meilleure Productivité",
    "7 Critères déterminants et 10 signaux d'alerte pour choisir le bon conjoint",
    "7 alertes avant que ton couples ne se brise",
    "7 secrets pour sauver son mariage pour monsieur",
    "7. secrets pour sauver son mariage pour madame",
    "8 jours pour expérimenter la surabondance de la grace",
    "9 portes d'accès au diable à fermer impérativement - Volume 1",
    "9 portes d'accès au diable à fermer impérativement - Volume 2",
    "À l'école de l'humilité, de l'obéissance",
    "Abraham et Sarah",
    "Accepter de perdre pour mieux réussir",
    "Accomplissez la vison de Dieu",
    "Achab et Jézabel, Pilate et Procla",
    "Achat",
    "ACTION DE GRACE",
    "Adam et Ève, quelles leçons tirons n",
    "Administration de l'Église",
    "Aimez-la et Respectez-le",
    "Aller en profondeur et faire plus",
    "Ambassadeur pour Christ",
    "Amplifiez votre Ministère par les miracles & les manifestations du Saint-Esprit",
    "Anagkazo",
    "Apprendre à gérer les conflits",
    "Arrêtons de juger pour mieux réussir",
    "Assoiffé de Dieu",
    "Assurance, amour, vie",
    "Avant de dire oui pour le mariage",
    "BATISSEURS",
    "Beaucoup sont appelés",
    "BEMA - Jugement et Justice",
    "Bible bilingue ESV / Colombe",
    "Bible d'étude Version du Semeur (Couverture bleue avec zip)",
    "Bible d'étude Version du Semeur (Couverture Noire)",
    "BIBLE D'ÉTUDE SEMEUR 2015 souple brune, marron, tranche blanche",
    "Bible Esprit et Vie Noir et argent",
    "BIBLE ESPRIT ET VIE - EDITION NUIT PU BLEU",
    "Bible Femmes à son écoute souple corail",
    "Bible Louis Segond 1910, Gros caractères format compact, similicuir, brun foncé avec onglets",
    "Bible Louis Segond 1910, Gros caractères format compact, similicuir, marine, onglets",
    "Bible Louis Segond 1910, Gros caractères, noir avec fermeture éclair",
    "Bible version francais courant",
    "Bible version LS 1910",
    "BIBLE SEGOND 1910, vinyle souple, Elegance bleu marine, moyenne, texte confort",
    "Bien choisir votre futur conjoint",
    "Bien penser pour mieux réussir",
    "Blessé par un père blessé",
    "Breaking free from strongholds",
    "Cahier d'etude",
    "CAMP DE REVEIL ADN",
    "Car on donnera à celui qui a et à celui qui n'a pas on ôtera même ce qu'il a",
    "Carnet des couples",
    "Casse-tête Jésus aime tout le monde",
    "Ce que signifie Devenir Berger",
    "Ce que signifie être prudent comme un serpent",
    "Ceux qui font semblant",
    "Ceux qui oublient",
    "Ceux qui sont des fils dangereux",
    "Ceux qui sont ignorants",
    "Ceux qui sont offensés",
    "Ceux qui sont orgueilleux",
    "Ceux qui vous accusent",
    "Ceux qui vous honorent",
    "Ceux qui vous quittent",
    "Chambre Occupation Double Avec Transport",
    "Chambre Occupation Double Avec Transport ancien",
    "Chambre Occupation Double Sans Transport",
    "Chambre Occupation Double Sans Transport ancien",
    "Chambre Occupation Quadruple Avec Transport",
    "Chambre Occupation Quadruple Avec Transport ancien",
    "Chambre Occupation Quadruple Sans Transport",
    "Chambre Occupation Quadruple Sans Transport 2",
    "Chambre Occupation Quadruple Sans Transport ancien",
    "Chambre Occupation Simple",
    "Chef de famille jusqu'au bout",
    "Choisi et établi pour porter du fruit qui démeure",
    "Christ en nous",
    "Comment accompagner nos ados",
    "Comment aider nos enfants à mieux réussir",
    "Comment bien organiser sa vie - Message pour la vie 1",
    "Comment bien remplir votre Ministère?",
    "Comment gagner votre conjoint à Christ",
    "Comment naître de nouveau et éviter l'enfer",
    "Comment neutraliser les malédictions?",
    "Comment prêcher le Salut?",
    "Comment prier?",
    "Comment puis-je dire merci?",
    "Comment témoigner de Christ",
    "Comment vivre le répos dans votre couple",
    "Comment votre langue peut détruire votre vie",
    "Comment vous pouvez avoir un temps de recueillement efficace avec Dieu tous les jours?",
    "Comment vous pouvez devenir un chrétien fort?",
    "Communiquer pour construire",
    "Complices et Audacieux",
    "Connaissez vos ennemis invisibles",
    "Conquérant Vainqueur",
    "Construire un mariage heureux et durable",
    "Couronnés de succès: 6 clés divines pour réussir en affaires et dans vos projets personnels",
    "Croire, confesser et pratiquer la parole de Dieu",
    "Croissance de l'église, c'est possible!",
    "Célibataire? Plus pour longtemps!",
    "Dangers spirituels",
    "David et Goliath",
    "destiné à Briller Vol 3 # CLés Maitresses pour parvenir au. sommet et briller",
    "Destiné à briller Vol 1 la guerre des étoiles",
    "Destiné à briller Vol2 LE PROTOCOE DIVIN POUR BRILLE",
    "Destinés à briller",
    "Dieu, le premier à reconnaître des droits aux femmes",
    "DIME",
    "Dites-leur: 120 raisons pour lesquelles vous devez être gagneur d'âmes",
    "Du Ghetto au Barreau",
    "Déjouer Les Pièges Du Manque De Pardon : Les Clés Qui Ont Transformé Ma Vie",
    "Eduquons adéquatement nos enfants",
    "ELGO",
    "ELGO-LIBRAIRIE",
    "En Christ toutes choses sont devenues nouvelles",
    "En temps de crise … Le juste que fera t-il?",
    "Entrez dans votre bénédiction",
    "EJP",
    "Espoir",
    "Est-il possible ?",
    "Étapes menant à l'onction",
    "Épouse du pasteur, Mère de l'église locale",
    "Éthique ministérielle",
    "Être fructueux",
    "ESTHER L'EPOUSE REVEILEE PAR LA TENTATIVE",
    "Fais Ta Part",
    "Faites de vous les sauveurs des hommes",
    "Femme , tu as une destinée",
    "FINAL TOUR",
    "Gagnez la guerre des pensées",
    "Gouverner avec la personne la plus importante",
    "Guide d'accompagnement sortir de la Blessure du Père",
    "HISTOIRE BIBLIQUE DE LA BIBLE APP",
    "IDENTITE: Confessions de Vie et Déclarations Prophétiques",
    "Il est temps que je rachete le. temps",
    "Il est temps que ma. vie soit productive",
    "Ils sont partis",
    "impacter et diriger en milieu hostile VOL1",
    "Implantation d'églises",
    "Incomparable",
    "INSATISFAIT : Entre dans la dimension de l'épouse",
    "Inébranlable",
    "Isaac et Rebecca",
    "J'ai donné ma vie à Jésus",
    "Je décide de restaurer mon âme",
    "Je me prépare pour dire oui",
    "Je quitte la prison de la peur",
    "Je suis un ambassadeur pour Christ",
    "Je suis une personne heureuse",
    "L'amour: la loi royale",
    "L'arbre et votre Ministère",
    "L'arrivée des enfants",
    "La bataille de l'image",
    "La belle, la bête & le Pasteur",
    "La Bible junior",
    "La Bible ne parle pas d'une réligion, mais d'un gouvernement",
    "La Bible Version du Semeur (couverture dure bleue)",
    "La Bible Version du Semeur (noir zip)",
    "La bonne semence 2025",
    "La ceinture de la vérité - Vol 2 - Les 7 armes de Dieu qui font fuir le diable",
    "La chair",
    "La communion du Saint-Esprit",
    "La cuirasse de la justice - Vol 3 - Les 7 armes de Dieu qui font fuir le diable",
    "La Direction Divine et les Lois Spirituelles du succès",
    "La double méga-église missionnaire",
    "La faveur divine",
    "La femme sage qui batit son foyer",
    "La foire aux questions de la vie chrétienne",
    "La Guerre Des Etoiles",
    "La guerre spirituelle - Vol 1 - Les 7 armes de Dieu qui font fuir le diable",
    "La Langue",
    "La lumière pour mieux gérer ses priorités - Message pour la vie 2",
    "La mentalité des personnes exceptionnelles",
    "La Méga-Église - Comment faire grandir votre église?",
    "La Parole de ma patience",
    "La perception divine: Bien percevoir pour mieux réussir",
    "La priere qui n'échoue jamais",
    "La prière",
    "La prière (Anglais)",
    "La prière exaucée à 100%",
    "La prédestination",
    "La préparation de l'évangile",
    "La présence de Dieu",
    "La prospérité du royaume des cieux",
    "La puissance de la vision",
    "La puissance d'une jeunesse choisie et révélée par Dieu",
    "La reine Esther",
    "La rétrogradation",
    "La révolution des amoureux de Dieu",
    "La sagesse est la chose principale pour votre Ministère",
    "La Sainte Bible",
    "La Sainte Bible (Avec œillet)",
    "La Sainte Bible (petit format bleu avec fermeture zip)",
    "La Sainte Bible (Sans œillet)",
    "La Sainte Bible (SOUPLE, MOYENNE, NOIR, TEXTE CONFORT)",
    "La Sainte Bible LS Violet avec oeillets",
    "La Sainte Bible Mini couverture zip bordeau or",
    "La Sainte Bible Mini couverture zip bleu argent",
    "La Sainte Bible, Souple DUO BLEU",
    "La Sainte Bible LS vinyle souple, Elegance bleu marine, moyenne, texte confort",
    "La Sainte Bible, Souple DUO BLEU",
    "La stérilité dans le Ministère",
    "La stérilité vaincue",
    "La vie abondante en Christ",
    "La voix que tu entends tracera ta voie",
    "LA BIBLE - Amis pour toujours",
    "LA NOUVELLE BIBLE SEGOND (NBS) - Édition d'étude",
    "LA SAINTE BIBLE BORDEAU RIGIDE",
    "LA SAINTE BIBLE COLOMBE NOIRE SEMI RIG",
    "LA SAINTE Bible LS1910, Gros caractères format compact, bordeaux avec onglets et fermeture éclair",
    "LA SAINTE Bible LS1910,marine, onglets",
    "Laissez-vous conduire par l'onction",
    "Laïkos",
    "Le baiser du ciel",
    "Le combat de la destinée",
    "Le combat de la Foi",
    "Le Coeur Du Chantre Vol 1",
    "Le Cœur du Chantre: Volume 1",
    "Le don de gouverner",
    "Le guide pratique de priere Tome I",
    "Le guide pratique de priere Tome II",
    "Le mariage modèle",
    "Le Mystère de la nouvelle création",
    "Le mystère de la foi",
    "Le mystère de la foi révélée",
    "Le nom de Jésus",
    "Le Nouveau Testament en bandes dessinées",
    "Le pardon rendu facile",
    "Le pouvoir a changé de main",
    "Le privilège",
    "Le regime de la grace",
    "Le réveil du premier amour",
    "Le Saint Esprit mon associé",
    "Le Saint-Esprit notre parakletos",
    "Le sang de Jésus - La puissance du Sang",
    "Le trio gagnant d'une vie équilibrée et épanouie",
    "Le Vrai disciple de Jesus-Christ",
    "les 4M de la femme de Dieu",
    "Les autres ...",
    "Les blessures de l'âme et le temps",
    "Les démons et comment les affronter?",
    "Les dix principales erreurs que commettent les Pasteurs",
    "Les douces influence de l'onction",
    "Les étapes menant à la Présence de Dieu",
    "Les facteurs déterminants",
    "Les gouteurs et les participants",
    "Les héritiers violents qui s'emparentdu Royaume des Cieux",
    "Les Laïcs & le Ministère",
    "Les pleurs et les grincements",
    "Les secrets de l'expansion",
    "Les secrets de la foi",
    "Les secrets de la paix intérieure",
    "Les secrets de la victoire",
    "Lis ta Bible",
    "Livre de prière FLOW",
    "Loyauté et déloyauté",
    "Ma fille, tu peux y arriver",
    "Ma première Bible à portée de main",
    "Maintenant ça suffit (Dur)",
    "Maintenant ça suffit (Souple)",
    "Maintenant ça suffit. Il faut que ça change",
    "Maintenant que vous êtes fiancés",
    "Manteau invisible",
    "Manuel de Cheminement Sortir de la Blessure du pere",
    "Manuel de mémorisation de la Bible",
    "Mariage sur mesure",
    "Marie, Celle qui a accepté de tout perdre",
    "Marie-toi bien",
    "Maîtriser l'art de l'épargne",
    "Mes premiers pas avec la Bible",
    "MINI LA SAINTE BIBLE AVEC POCHETTE",
    "Moise et Sephora",
    "Mon ambition c'est Christ",
    "Mon Dieu pourvoit à tous mes besoins",
    "Mon juste vivra par la Foi",
    "Mon journal de productivité",
    "Mon père mon mari",
    "MTL-LIBRAIRIE",
    "Naitre dans le Royaume (10-14 ans)",
    "Naitre dans le Royaume (3-9 ans)",
    "Naître et grandir dans le Royaume des Cieux",
    "Nathan, Koré, Abiram",
    "Ne brade pas ton futur",
    "Ne peux-tu pas faire un peu plus?",
    "Neutraliser l'ennemi public numéro 1 (Anglais)",
    "Neutraliser l'ennemi public numéro 1 , la chair Vol 2",
    "Neutralize the n°1 public enemy : the flesh",
    "Nommer-le! Réclamez-le!! Prenez-le!!!",
    "Offrande retraite ADN 2026",
    "Offrande Women's Camp 2026",
    "Offrandes Camps Embrasement",
    "OFFRANDES",
    "Osez être vous-même",
    "OTTAWA",
    "Pas un novice",
    "Perdre, souffrir, sacrifier, mourir",
    "Possédez les clefs de votre règne",
    "PREMICES",
    "Pourquoi le mariage est autant combattu?",
    "Pourquoi les chrétiens qui ne paient pas la dime deviennent pauvre …",
    "Pourquoi peu sont élus?",
    "Précieuse comme un diamant",
    "Prêt à 20 ans",
    "Prions ensemble (vol. 1)",
    "Prions ensemble (vol. 2)",
    "Prions ensemble (vol. 3)",
    "Protocole Divin pour briller",
    "Quel genre d'homme etes-vous dans votre maison",
    "Quel genre de femme suis-je dans ma maison",
    "Que ton règne vienne",
    "Qu'est ce qui vous empêche de vous marier",
    "Quand le rejet devient une porte",
    "Ranimez-vous",
    "Recevoir l'onction",
    "Restaurer la conscience de Dieu dans la famille",
    "Restaurer l'âme blessée",
    "Retraite EJP ADN 2026 formule sans hébergement et sans repas",
    "Reussir et régner dans le royaume",
    "Réservé aux parents",
    "Révolution Ados",
    "Révolution Ados",
    "Règles pour le Ministère à plein temps",
    "Règles pour le travail au sein de l'Église",
    "RUTH CELLE QUE L'ON N'ATTENDAIT PAS",
    "Seigneur, je sais que Tu as besoin de quelqu'un",
    "Seminaire - PST LILLIANE SANOGO",
    "SEMENCE",
    "Sept grands principes",
    "Servir selon le cœur de Dieu",
    "Si vous aimez le Seigneur",
    "Silence coupable",
    "Sois fidèle jusqu'à la mort",
    "Son pouvoir, Sa présence",
    "Sortir des Prisons Intérieures",
    "Sortir des prisons (Anglais)",
    "Sortir des prisons intérieures",
    "STOP, ma vie doit compter",
    "Suis-je bon à rien?",
    "Ta Volonté",
    "Tentez de grandes choses pour Dieu",
    "The prayer that never fails",
    "test",
    "Tout par la prière, rien sans la prière",
    "Tout pour une famille harmonieuse",
    "Transport Retraite EJP ADN 2026",
    "Travaillez pour être bénis. Ne travaillez pas pour être riche",
    "Transformez votre Ministère pastoral",
    "Un bon Général - La science du Leadership",
    "Une vie restaurée pour une destinée glorieuse",
    "Université des tueurs de géants",
    "VictorIAVILLE",
    "Voir Christ tel il est",
    "Voir et entendre",
    "Vivez l'extraordinaire divin Vol 1",
    "Vivre et apprécier son Célibat",
    "VOEUX",
    "Vol 1_Les 10 handicaps du leader qui freinent la croissance du ministère",
    "Vol 2_Les 10 handicaps du leader qui freinent la croissance du ministère",
    "Vos premières années de vie conjugale",
    "Votre cœur face",
    "Votre cœur face à la parole",
    "Vous avez rompu avant l'étape du mariage",
    "Vous pensez mariage? Comment faire le bon choix?",
    "Zacharie",
    "Zakarie et Elisabeth - Joseph et Marie",
    "annuler Chambre Occupation Quadruple Avec Transport",
    "annuler Chambre Occupation Quadruple Sans Transport",
    "annuler Chambre Occupation Simple",
    "annulerChambre Occupation Double Avec Transport",
    "annulerChambre Occupation Double Sans Transport",
    "chambre occupation Simple Sarah",
    "la foi persévérante",
]


# =========================================================
# CONFIGURATION
# =========================================================

NOM_FICHIER = "ReportTemplate.xlsx"

COLONNE_1 = "Librairie"
COLONNE_2 = "Articles"
COLONNE_3 = "StockRestant"

PREMIERE_LIGNE_DONNEES = 2
DERNIERE_LIGNE = 1000


# =========================================================
# CRÉATION DU CLASSEUR
# =========================================================

wb = Workbook()

ws = wb.active
ws.title = "Données"

# Feuille contenant les listes sources
ws_listes = wb.create_sheet("Listes")


# =========================================================
# EN-TÊTES
# =========================================================

ws["A1"] = COLONNE_1
ws["B1"] = COLONNE_2
ws["C1"] = COLONNE_3


# =========================================================
# LISTES SOURCES
# =========================================================

# Liste du premier type
for ligne, valeur in enumerate(LIBRAIRIES, start=1):
    ws_listes.cell(row=ligne, column=1, value=valeur)

# Liste du deuxième type
for ligne, valeur in enumerate(ARTICLES, start=1):
    ws_listes.cell(row=ligne, column=2, value=valeur)


# =========================================================
# MENUS DÉROULANTS
# =========================================================

# Menu déroulant pour la colonne A
validation_type_1 = DataValidation(
    type="list",
    formula1=f"=Listes!$A$1:$A${len(LIBRAIRIES)}",
    allow_blank=True
)

# Menu déroulant pour la colonne B
validation_type_2 = DataValidation(
    type="list",
    formula1=f"=Listes!$B$1:$B${len(ARTICLES)}",
    allow_blank=True
)


# Ajouter les validations à la feuille
ws.add_data_validation(validation_type_1)
ws.add_data_validation(validation_type_2)


# Appliquer les menus déroulants aux colonnes
validation_type_1.add(
    f"A{PREMIERE_LIGNE_DONNEES}:A{DERNIERE_LIGNE}"
)

validation_type_2.add(
    f"B{PREMIERE_LIGNE_DONNEES}:B{DERNIERE_LIGNE}"
)


# =========================================================
# MISE EN FORME MINIMALE
# =========================================================

ws.column_dimensions["A"].width = 25
ws.column_dimensions["B"].width = 25

ws_listes.column_dimensions["A"].width = 25
ws_listes.column_dimensions["B"].width = 25


# =========================================================
# SAUVEGARDE
# =========================================================

wb.save(NOM_FICHIER)

print(f"Fichier créé : {NOM_FICHIER}")